"""
Document text extraction utility supporting TXT, MD, PDF, and DOCX files.
Performs text sanitation and cleaning.
"""

import io
import re
from pathlib import Path
from typing import Union


class DocumentExtractor:
    """Extracts raw text from various document formats with robust error handling."""

    @classmethod
    def extract_text(cls, file_content: bytes, filename: str) -> str:
        """
        Extract text from file bytes based on file extension.
        """
        ext = Path(filename).suffix.lower()

        if ext in [".txt", ".md"]:
            return cls._extract_from_text(file_content)
        elif ext == ".pdf":
            return cls._extract_from_pdf(file_content)
        elif ext in [".docx", ".doc"]:
            return cls._extract_from_docx(file_content)
        else:
            raise ValueError(f"Unsupported file format '{ext}'. Supported formats: .txt, .md, .pdf, .docx")

    @classmethod
    def _extract_from_text(cls, content: bytes) -> str:
        """Decode plain text or markdown with encoding fallbacks."""
        for encoding in ["utf-8", "utf-8-sig", "latin-1", "cp1252"]:
            try:
                text = content.decode(encoding)
                return cls.clean_text(text)
            except UnicodeDecodeError:
                continue
        raise ValueError("Failed to decode text file. Ensure it is a valid text encoding.")

    @classmethod
    def _extract_from_pdf(cls, content: bytes) -> str:
        """Extract text from PDF pages using pypdf."""
        import pypdf
        try:
            reader = pypdf.PdfReader(io.BytesIO(content))
            pages_text = []
            for i, page in enumerate(reader.pages):
                extracted = page.extract_text()
                if extracted:
                    pages_text.append(extracted)
            if not pages_text:
                raise ValueError("PDF does not contain any readable extractable text.")
            return cls.clean_text("\n\n".join(pages_text))
        except Exception as e:
            if isinstance(e, ValueError):
                raise
            raise ValueError(f"Failed to extract text from PDF: {str(e)}")

    @classmethod
    def _extract_from_docx(cls, content: bytes) -> str:
        """Extract text from Word .docx file using python-docx."""
        import docx
        try:
            doc = docx.Document(io.BytesIO(content))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                    if row_text:
                        paragraphs.append(row_text)
            if not paragraphs:
                raise ValueError("DOCX document is empty.")
            return cls.clean_text("\n\n".join(paragraphs))
        except Exception as e:
            if isinstance(e, ValueError):
                raise
            raise ValueError(f"Failed to extract text from DOCX: {str(e)}")

    @classmethod
    def clean_text(cls, text: str) -> str:
        """Sanitize whitespace, remove null characters, and normalize line breaks."""
        text = text.replace("\x00", "")
        text = re.sub(r"\r\n", "\n", text)
        text = re.sub(r"\n{3,}", "\n\n", text)
        text = re.sub(r"[ \t]+", " ", text)
        return text.strip()

